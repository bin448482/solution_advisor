from pathlib import Path
from typing import List

from pptx import Presentation

from src.models import SlideText


def extract_text(pptx_path: Path) -> List[SlideText]:
    """Extract title, text frames, and notes from PPTX."""
    presentation = Presentation(str(pptx_path))
    slides: List[SlideText] = []
    for idx, slide in enumerate(presentation.slides, start=1):
        title = slide.shapes.title.text if slide.shapes.title else None
        body_parts = []
        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue
            text = shape.text.strip()
            if text:
                body_parts.append(text)
        text_content = "\n".join(body_parts)

        notes_text = None
        if getattr(slide, "has_notes_slide", False) and slide.notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = slide.notes_slide.notes_text_frame.text

        slides.append(
            SlideText(
                slide_no=idx,
                title=title,
                text_content=text_content,
                notes=notes_text,
            )
        )
    return slides
